# WorkoutLoggerApp.py
import streamlit as st
import pandas as pd
from openpyxl import load_workbook

EXCEL_FILE = "Workout_Progress_Tracker.xlsx"
SHEET_NAME = "Logger"

st.title("Workout Logger 📓")

# Load existing data
try:
    df_existing = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME)
except FileNotFoundError:
    st.error(f"Could not find {EXCEL_FILE}. Make sure it's in the same folder.")
    st.stop()

st.write("Current log (read-only view):")
st.dataframe(df_existing)

# Streamlit form
with st.form("log_form"):
    date = st.date_input("Date")
    day = st.selectbox("Day (Pull/Push/Legs)", ["Pull", "Push", "Legs"])
    exercise = st.text_input("Exercise")
    set_number = st.number_input("Set #", min_value=1, value=1)
    reps = st.number_input("Reps", min_value=1, value=10)
    weight_kg = st.number_input("Weight (kg)", min_value=0.0, step=0.5)
    weight_lbs = st.number_input("Weight (lb)", min_value=0.0, step=1.0)
    notes = st.text_area("Notes")
    submitted = st.form_submit_button("Add Workout")

if submitted:
    # Convert weight
    weight_converted = weight_kg if weight_kg > 0 else round(weight_lbs / 2.20462, 2)

    new_row = {
        "Date": str(date),
        "Day (Pull/Push/Legs)": day,
        "Exercise": exercise,
        "Set #": set_number,
        "Reps": reps,
        "Weight (kg)": weight_kg,
        "Weight (lb)": weight_lbs,
        "Weight (kg converted)": weight_converted,
        "Notes": notes
    }

    # Always start appending after row 7
    book = load_workbook(EXCEL_FILE)
    sheet = book[SHEET_NAME]

    next_row = 8
    while sheet.cell(row=next_row, column=1).value not in (None, ""):
        next_row += 1

    for idx, value in enumerate(new_row.values(), start=1):
        sheet.cell(row=next_row, column=idx).value = value

    book.save(EXCEL_FILE)
    st.success(f"Workout added successfully to row {next_row}!")

    # Refresh view
    df_existing = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_NAME)
    st.dataframe(df_existing)

